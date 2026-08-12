"""
Revenue Report Service — Admin Dashboard

Read-only service for the dedicated Revenue Report panel.
Business rule: revenue is defined only from orders with status='completed'.
Revenue is calculated at OrderItem level as quantity * price.
"""

from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Dict, List, Optional, Tuple

from dateutil.relativedelta import relativedelta
from django.core.paginator import Paginator
from django.db.models import Count, F, Q, Sum
from django.db.models.functions import TruncDay, TruncMonth, TruncQuarter, TruncYear
from django.conf import settings
from django.utils import timezone

from orders.models import Order, OrderItem
from products.models import Brand, Category, Product


ORDER_STATUSES = ['pending', 'processing', 'completed', 'cancelled']
PAYMENT_METHODS = ['cod', 'bank_transfer', 'momo', 'vnpay']


@dataclass
class RevenueFilters:
    period: str = 'day'
    date_from: Optional[date] = None
    date_to: Optional[date] = None
    top_limit: int = 10
    brand_ids: List[int] = field(default_factory=list)
    category_ids: List[int] = field(default_factory=list)
    payment_methods: List[str] = field(default_factory=list)
    order_statuses: List[str] = field(default_factory=list)


def _normalize_filters(filters: RevenueFilters) -> RevenueFilters:
    today = timezone.now().date()
    if filters.period not in ('day', 'month', 'quarter', 'year'):
        filters.period = 'day'
    if filters.date_from is None or filters.date_to is None:
        if filters.period == 'day':
            filters.date_to = today
            filters.date_from = today - timedelta(days=29)
        elif filters.period == 'month':
            filters.date_to = today
            filters.date_from = (today - relativedelta(months=11)).replace(day=1)
        elif filters.period == 'quarter':
            current_q = (today.month - 1) // 3
            start_month = current_q * 3 + 1
            current_quarter_start = today.replace(month=start_month, day=1)
            filters.date_from = current_quarter_start - relativedelta(months=9)
            filters.date_to = current_quarter_start + relativedelta(months=3) - timedelta(days=1)
        elif filters.period == 'year':
            filters.date_to = today.replace(month=12, day=31)
            filters.date_from = today.replace(year=today.year - 4, month=1, day=1)
    return filters


def _previous_period_filters(filters: RevenueFilters) -> RevenueFilters:
    """Return filters covering the same-length period immediately before date_from."""
    delta_days = (filters.date_to - filters.date_from).days
    prev_to = filters.date_from - timedelta(days=1)
    prev_from = prev_to - timedelta(days=delta_days)
    return RevenueFilters(
        period=filters.period,
        date_from=prev_from,
        date_to=prev_to,
        top_limit=filters.top_limit,
        brand_ids=list(filters.brand_ids),
        category_ids=list(filters.category_ids),
        payment_methods=list(filters.payment_methods),
    )


def _revenue_expression():
    return F('quantity') * F('price')


def _aggregate_revenue(qs):
    agg = qs.aggregate(
        revenue=Sum(_revenue_expression()),
        orders=Count('order', distinct=True),
    )
    revenue = agg['revenue'] or Decimal('0')
    orders = agg['orders'] or 0
    return revenue, orders


def _format_aov(revenue: Decimal, orders: int) -> Decimal:
    if orders == 0:
        return Decimal('0')
    return (revenue / orders).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def _build_item_queryset(filters: RevenueFilters, status: str = 'completed'):
    """Build OrderItem queryset for a given order status with all filters applied."""
    qs = OrderItem.objects.select_related('product', 'product__brand', 'product__category')
    order_filter = Q(order__status=status)
    if filters.date_from:
        order_filter &= Q(order__created_at__date__gte=filters.date_from)
    if filters.date_to:
        order_filter &= Q(order__created_at__date__lte=filters.date_to)
    if filters.payment_methods:
        order_filter &= Q(order__payment_method__in=filters.payment_methods)
    if filters.brand_ids:
        order_filter &= Q(product__brand_id__in=filters.brand_ids)
    if filters.category_ids:
        order_filter &= Q(product__category_id__in=filters.category_ids)
    return qs.filter(order_filter)


def _build_order_queryset(filters: RevenueFilters, status: str):
    """Build Order queryset for a given status with all filters applied (via items for brand/category)."""
    qs = Order.objects.all()
    order_filter = Q(status=status)
    if filters.date_from:
        order_filter &= Q(created_at__date__gte=filters.date_from)
    if filters.date_to:
        order_filter &= Q(created_at__date__lte=filters.date_to)
    if filters.payment_methods:
        order_filter &= Q(payment_method__in=filters.payment_methods)
    qs = qs.filter(order_filter)
    if filters.brand_ids:
        qs = qs.filter(items__product__brand_id__in=filters.brand_ids).distinct()
    if filters.category_ids:
        qs = qs.filter(items__product__category_id__in=filters.category_ids).distinct()
    return qs


def _period_label_for_key(period: str, key) -> str:
    if period == 'day':
        return key.strftime('%d/%m/%Y')
    if period == 'month':
        return f'{key.month:02d}/{key.year}'
    if period == 'quarter':
        return f'Q{key[1]}-{key[0]}'
    if period == 'year':
        return str(key)
    return str(key)


def _period_to_key(period: str, dt) -> Any:
    if period == 'day':
        return dt.date()
    if period == 'month':
        return date(dt.year, dt.month, 1)
    if period == 'quarter':
        return (dt.year, (dt.month - 1) // 3 + 1)
    if period == 'year':
        return dt.year
    return dt


def _generate_period_keys(period: str, date_from: date, date_to: date) -> List[Any]:
    keys = []
    current = date_from
    if period == 'day':
        while current <= date_to:
            keys.append(current)
            current += timedelta(days=1)
    elif period == 'month':
        current = current.replace(day=1)
        while current <= date_to:
            keys.append(date(current.year, current.month, 1))
            current += relativedelta(months=1)
    elif period == 'quarter':
        q = (current.month - 1) // 3
        start_month = q * 3 + 1
        current = current.replace(month=start_month, day=1)
        while current <= date_to:
            q = (current.month - 1) // 3 + 1
            keys.append((current.year, q))
            current += relativedelta(months=3)
    elif period == 'year':
        current = current.replace(month=1, day=1)
        while current <= date_to:
            keys.append(current.year)
            current = current.replace(year=current.year + 1)
    return keys


def _trunc_for_period(period: str):
    mapping = {
        'day': TruncDay('order__created_at'),
        'month': TruncMonth('order__created_at'),
        'quarter': TruncQuarter('order__created_at'),
        'year': TruncYear('order__created_at'),
    }
    return mapping[period]


def get_revenue_kpi(filters: RevenueFilters) -> Dict[str, Any]:
    filters = _normalize_filters(filters)
    revenue, orders = _aggregate_revenue(_build_item_queryset(filters))
    prev_filters = _previous_period_filters(filters)
    prev_revenue, _ = _aggregate_revenue(_build_item_queryset(prev_filters))
    growth_pct = None
    if prev_revenue > 0:
        growth_pct = round(float((revenue - prev_revenue) / prev_revenue * 100), 1)
    return {
        'revenue': revenue,
        'orders': orders,
        'aov': _format_aov(revenue, orders),
        'previous_revenue': prev_revenue,
        'growth_pct': growth_pct,
        'date_from': filters.date_from,
        'date_to': filters.date_to,
        'previous_date_from': prev_filters.date_from,
        'previous_date_to': prev_filters.date_to,
    }


def get_revenue_trend(filters: RevenueFilters) -> List[Dict[str, Any]]:
    filters = _normalize_filters(filters)
    trunc = _trunc_for_period(filters.period)
    qs = _build_item_queryset(filters).annotate(period=trunc)
    rows = qs.values('period').annotate(
        revenue=Sum(_revenue_expression()),
        orders=Count('order', distinct=True),
    ).order_by('period')
    rows_map = {}
    for row in rows:
        key = _period_to_key(filters.period, row['period'])
        rows_map[key] = row
    keys = _generate_period_keys(filters.period, filters.date_from, filters.date_to)
    result = []
    for key in keys:
        row = rows_map.get(key)
        rev = row['revenue'] if row else Decimal('0')
        ord_count = row['orders'] if row else 0
        result.append({
            'period_label': _period_label_for_key(filters.period, key),
            'period_key': key,
            'revenue': rev,
            'orders': ord_count,
            'aov': _format_aov(rev, ord_count),
        })
    return result


def get_revenue_comparison(filters: RevenueFilters) -> Dict[str, Any]:
    filters = _normalize_filters(filters)
    current_revenue, current_orders = _aggregate_revenue(_build_item_queryset(filters))
    prev_filters = _previous_period_filters(filters)
    prev_revenue, prev_orders = _aggregate_revenue(_build_item_queryset(prev_filters))
    growth_pct = None
    if prev_revenue > 0:
        growth_pct = round(float((current_revenue - prev_revenue) / prev_revenue * 100), 1)
    return {
        'current': {
            'revenue': current_revenue,
            'orders': current_orders,
            'aov': _format_aov(current_revenue, current_orders),
            'date_from': filters.date_from,
            'date_to': filters.date_to,
        },
        'previous': {
            'revenue': prev_revenue,
            'orders': prev_orders,
            'aov': _format_aov(prev_revenue, prev_orders),
            'date_from': prev_filters.date_from,
            'date_to': prev_filters.date_to,
        },
        'growth_pct': growth_pct,
    }


def get_revenue_breakdown(
    filters: RevenueFilters,
    group_by: str,
    limit: Optional[int] = None,
) -> List[Dict[str, Any]]:
    filters = _normalize_filters(filters)
    qs = _build_item_queryset(filters)
    if group_by == 'category':
        qs = qs.values('product__category__name', 'product__category__id').annotate(
            revenue=Sum(_revenue_expression()),
            quantity=Sum('quantity'),
            orders=Count('order', distinct=True),
        ).order_by('-revenue')
    elif group_by == 'brand':
        qs = qs.values('product__brand__name', 'product__brand__id').annotate(
            revenue=Sum(_revenue_expression()),
            quantity=Sum('quantity'),
            orders=Count('order', distinct=True),
        ).order_by('-revenue')
    elif group_by == 'product':
        qs = qs.values(
            'product__id', 'product__name', 'product__brand__name', 'product__category__name'
        ).annotate(
            revenue=Sum(_revenue_expression()),
            quantity=Sum('quantity'),
            orders=Count('order', distinct=True),
        ).order_by('-revenue')
    else:
        return []
    result = []
    for idx, row in enumerate(qs):
        if limit is not None and idx >= limit:
            break
        revenue = row['revenue'] or Decimal('0')
        quantity = row['quantity'] or 0
        result.append({
            'name': (
                row.get('product__name')
                or row.get('product__category__name')
                or row.get('product__brand__name')
                or '—'
            ),
            'revenue': revenue,
            'quantity': quantity,
            'orders': row['orders'] or 0,
            'avg_price': _format_aov(revenue, quantity),
        })
    return result


def get_top_revenue_products(filters: RevenueFilters, limit: int = 10) -> List[Dict[str, Any]]:
    return get_revenue_breakdown(filters, group_by='product', limit=limit)


def get_revenue_details(filters: RevenueFilters) -> List[Dict[str, Any]]:
    return get_revenue_trend(filters)


def get_cancel_impact(filters: RevenueFilters) -> Dict[str, Any]:
    filters = _normalize_filters(filters)
    completed_revenue, _ = _aggregate_revenue(_build_item_queryset(filters, status='completed'))
    cancelled_qs = _build_order_queryset(filters, status='cancelled')
    cancelled_count = cancelled_qs.count()
    cancelled_revenue = cancelled_qs.aggregate(total=Sum('total_price'))['total'] or Decimal('0')
    return {
        'completed_revenue': completed_revenue,
        'cancelled_orders': cancelled_count,
        'cancelled_revenue': cancelled_revenue,
    }


def get_revenue_report(filters: RevenueFilters, page: int = 1, per_page: int = 20) -> Dict[str, Any]:
    """Single entry point returning all Revenue Report sections."""
    filters = _normalize_filters(filters)
    trend = get_revenue_trend(filters)
    paginator = Paginator(trend, per_page)
    page_obj = paginator.get_page(page)
    return {
        'kpi': get_revenue_kpi(filters),
        'trend': trend,
        'comparison': get_revenue_comparison(filters),
        'breakdown_category': get_revenue_breakdown(filters, 'category', limit=filters.top_limit),
        'breakdown_category_all': get_revenue_breakdown(filters, 'category', limit=None),
        'breakdown_brand': get_revenue_breakdown(filters, 'brand', limit=filters.top_limit),
        'breakdown_brand_all': get_revenue_breakdown(filters, 'brand', limit=None),
        'breakdown_product': get_revenue_breakdown(filters, 'product', limit=filters.top_limit),
        'breakdown_product_all': get_revenue_breakdown(filters, 'product', limit=None),
        'top_products': get_top_revenue_products(filters, limit=filters.top_limit),
        'cancel': get_cancel_impact(filters),
        'page_obj': page_obj,
        'paginator': paginator,
    }


def _prepare_export_rows(filters: RevenueFilters) -> Tuple[List[str], List[List[Any]]]:
    filters = _normalize_filters(filters)
    trend = get_revenue_trend(filters)
    headers = ['Kỳ', 'Số đơn hoàn thành', 'Doanh thu', 'AOV']
    rows = []
    for row in trend:
        rows.append([
            row['period_label'],
            row['orders'],
            float(row['revenue']),
            float(row['aov']),
        ])
    return headers, rows


def export_revenue_xlsx(filters: RevenueFilters, request=None) -> bytes:
    import io
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    filters = _normalize_filters(filters)
    headers, rows = _prepare_export_rows(filters)

    store_name = getattr(settings, 'STORE_NAME', 'TechGear')
    period_map = {
        'day': 'Ngày',
        'month': 'Tháng',
        'quarter': 'Quý',
        'year': 'Năm',
    }
    period_text = period_map.get(filters.period, 'Ngày')
    date_from_str = filters.date_from.strftime('%d/%m/%Y') if filters.date_from else ''
    date_to_str = filters.date_to.strftime('%d/%m/%Y') if filters.date_to else ''
    export_datetime = timezone.now().strftime('%d/%m/%Y %H:%M')
    preparer = (
        request.user.name
        if request and request.user.is_authenticated and getattr(request.user, 'name', '')
        else 'Người lập'
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Revenue Report'

    col_count = len(headers)
    title_font = Font(bold=True, size=16)
    label_font = Font(bold=True, size=11)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='6366F1', end_color='6366F1', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center')
    right = Alignment(horizontal='right', vertical='center')
    thin = Side(border_style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def write_merged(row: int, value: Any, font: Font, align: Alignment = center):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
        cell = ws.cell(row=row, column=1, value=value)
        cell.font = font
        cell.alignment = align

    write_merged(1, store_name, Font(bold=True, size=12), center)
    write_merged(2, 'BÁO CÁO DOANH THU', title_font, center)
    write_merged(3, f'Từ ngày {date_from_str} đến ngày {date_to_str}', Font(size=11), center)
    write_merged(4, f'Theo: {period_text}', Font(size=11), center)
    write_merged(5, f'Ngày xuất báo cáo: {export_datetime}', Font(size=11), right)

    header_row = 7
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(rows, header_row + 1):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = right if col_idx > 1 else left
            cell.border = border

    signature_row = header_row + len(rows) + 2
    ws.cell(row=signature_row, column=1, value='Người lập').font = label_font
    ws.cell(row=signature_row, column=4, value='Người duyệt').font = label_font
    ws.cell(row=signature_row + 1, column=1, value=preparer)
    ws.cell(row=signature_row + 1, column=4, value='')
    ws.cell(row=signature_row + 3, column=1, value='(Ký, họ tên)')
    ws.cell(row=signature_row + 3, column=4, value='(Ký, họ tên)')

    for col_idx in range(1, col_count + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 24

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
